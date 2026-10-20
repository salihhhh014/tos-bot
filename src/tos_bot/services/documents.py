from __future__ import annotations

import io

from aiogram import Bot
from aiogram.types import Message
from docx import Document as DocxDocument
from pypdf import PdfReader

from tos_bot.config import Settings
from tos_bot.services.ocr import MIME as OCR_MIME
from tos_bot.services.ocr import ocr_document

ALLOWED_EXT = {".pdf", ".docx", ".txt", ".md", ".jpg", ".jpeg", ".png"}


async def extract_from_telegram_file(
    bot: Bot | None, msg: Message, cfg: Settings
) -> tuple[str, bool]:
    """Возвращает (текст, был_ли_ocr). Текста нет и OCR не помог — ValueError."""
    doc = msg.document
    if doc is None or bot is None:
        raise ValueError("нет файла")

    size = doc.file_size or 0
    if size > cfg.max_file_mb * 1024 * 1024:
        raise ValueError(f"файл больше {cfg.max_file_mb} МБ")

    name = (doc.file_name or "").lower()
    ext = next((e for e in ALLOWED_EXT if name.endswith(e)), None)
    if ext is None:
        raise ValueError("давай pdf, docx, txt или фото (jpg/png)")

    raw = await _download(bot, doc.file_id, cfg)
    return await extract_bytes(raw, ext, cfg)


async def extract_from_telegram_photo(
    bot: Bot | None, msg: Message, cfg: Settings
) -> tuple[str, bool]:
    """Скриншот, отправленный как фото (не файлом), сразу уходит в OCR."""
    if bot is None or not msg.photo:
        raise ValueError("нет фото")

    photo = msg.photo[-1]
    if (photo.file_size or 0) > cfg.max_file_mb * 1024 * 1024:
        raise ValueError(f"фото больше {cfg.max_file_mb} МБ")

    raw = await _download(bot, photo.file_id, cfg)
    return await extract_bytes(raw, ".jpg", cfg)


async def _download(bot: Bot, file_id: str, cfg: Settings) -> bytes:
    tg_file = await bot.get_file(file_id)
    buf = io.BytesIO()
    await bot.download_file(tg_file.file_path or "", destination=buf)
    raw = buf.getvalue()
    if len(raw) > cfg.max_file_mb * 1024 * 1024:
        raise ValueError(f"файл больше {cfg.max_file_mb} МБ")
    return raw


async def extract_bytes(raw: bytes, ext: str, cfg: Settings) -> tuple[str, bool]:
    if ext == ".pdf":
        text = _pdf_to_text(raw)
    elif ext == ".docx":
        text = _docx_to_text(raw)
    elif ext in OCR_MIME:
        text = ""  # картинки сразу в OCR, извлекать нечего
    else:
        text = _txt_to_text(raw)

    if len(text.strip()) >= 200:
        return text, False

    # текста почти нет — похоже на скан. пробуем распознать
    if ext not in OCR_MIME:
        raise ValueError("в файле почти нет текста")
    if not cfg.ocr_enabled:
        raise ValueError("похоже на скан без текстового слоя, а OCR выключен (OCR_ENABLED=1)")
    return await ocr_document(cfg, raw, ext), True


def _pdf_to_text(raw: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(raw))
        if len(reader.pages) > 100:
            raise ValueError("pdf слишком длинный (больше 100 страниц)")
        parts = [(p.extract_text() or "") for p in reader.pages]
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("pdf повреждён или защищён паролем") from exc
    return "\n".join(parts)


def _docx_to_text(raw: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(raw))
        # таблицы тоже читаем — в офертах деньги часто лежат там
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text for c in row.cells))
    except Exception as exc:
        raise ValueError("docx повреждён или это не Word-документ") from exc
    return "\n".join(lines)


def _txt_to_text(raw: bytes) -> str:
    # utf-8 с заменой — лучше кривые символы, чем падение
    return raw.decode("utf-8", errors="replace")
